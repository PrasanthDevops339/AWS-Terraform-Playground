"""
Script to generate the config_aggregator.py How It Works Word document.
Run: python generate_config_aggregator_doc.py
Output: config_aggregator_how_it_works.docx in the same directory
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os


def set_cell_shading(cell, color_hex):
    from lxml import etree
    shading = cell._element.get_or_add_tcPr()
    shading_elem = etree.SubElement(shading, qn('w:shd'))
    shading_elem.set(qn('w:fill'), color_hex)
    shading_elem.set(qn('w:val'), 'clear')


def add_styled_table(doc, headers, rows, col_widths=None, header_color='2E74B5'):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, header_color)
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(val)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if r % 2 == 1:
                set_cell_shading(cell, 'DEEAF6')
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(width)
    return table


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)
    return p


def main():
    doc = Document()

    # -------------------------------------------------------------------------
    # Title
    # -------------------------------------------------------------------------
    title = doc.add_heading('config_aggregator.py — How It Works', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('AWS Config Compliance Reporting Pipeline  |  ECS Fargate Scheduled Task')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x59, 0x56, 0x59)

    doc.add_paragraph()

    meta_headers = ['Field', 'Value']
    meta_rows = [
        ['Script', 'scripts/config_aggregator.py'],
        ['Runtime', 'AWS ECS Fargate (scheduled task)'],
        ['Purpose', 'Query AWS Config aggregator for non-compliant resources across all org accounts'],
        ['Output', 'One CSV file per account uploaded to S3'],
        ['Language', 'Python 3'],
    ]
    add_styled_table(doc, meta_headers, meta_rows, col_widths=[1.8, 4.7])

    # -------------------------------------------------------------------------
    # Section 1: Overview
    # -------------------------------------------------------------------------
    doc.add_heading('1. Overview', level=1)
    doc.add_paragraph(
        'config_aggregator.py is a compliance reporting pipeline. It runs as a scheduled '
        'ECS task and produces one CSV file per AWS account containing all non-compliant '
        'resources for that account, enriched with rule annotation details and account metadata.'
    )
    doc.add_paragraph('At a high level the script does five things in sequence:')
    for step in [
        '1. Read the active ingest policy from DynamoDB (what resource types and rules to query).',
        '2. Query the AWS Config aggregator for all NON_COMPLIANT resources matching the policy.',
        '3. For each resource: check exclusions (Suspended OU, 1.0 legacy accounts), '
           'retrieve rule annotations, and apply keyword filters.',
        '4. Write passing resources to an in-memory CSV buffer.',
        '5. Group results by account and upload one CSV per account to S3.',
    ]:
        doc.add_paragraph(step, style='List Number')

    # -------------------------------------------------------------------------
    # Section 2: Environment Variables
    # -------------------------------------------------------------------------
    doc.add_heading('2. Environment Variables', level=1)
    doc.add_paragraph(
        'All configuration is injected via environment variables at ECS task launch time. '
        'No configuration is hardcoded in the script.'
    )
    env_headers = ['Variable', 'Purpose', 'Example']
    env_rows = [
        ['AGGREGATOR_NAME', 'Name of the AWS Config multi-account aggregator', 'org-config-aggregator'],
        ['REGION', 'AWS region for all API calls', 'us-east-2'],
        ['TAGGING_BUCKET', 'S3 bucket where output CSVs are uploaded', 'my-compliance-reports'],
        ['BUCKET_PREFIX', 'S3 key prefix for organising output files', 'config-reports/daily'],
        ['POLICY_TABLE', 'DynamoDB table holding the active policy/rule configuration', 'operations-prod-policies'],
        ['CLOUD_VERSION_TABLE', 'DynamoDB table listing 1.0 (legacy) accounts to exclude', 'operations-prod-cloud-versions'],
        ['ACCOUNT_ID_ALLOWLIST', 'Optional comma-separated list to restrict query scope', '111122223333,222233334444'],
        ['SUSPENDED_OU_CACHE_TTL_ENABLED', 'Whether the suspended OU cache expires (default true)', 'true'],
        ['SUSPENDED_OU_CACHE_TTL_SECONDS', 'Suspended OU cache lifetime in seconds (default 1800)', '1800'],
    ]
    add_styled_table(doc, env_headers, env_rows, col_widths=[2.2, 2.6, 1.7])

    # -------------------------------------------------------------------------
    # Section 3: Entry Point and Startup
    # -------------------------------------------------------------------------
    doc.add_heading('3. Script Entry Point and Startup', level=1)
    doc.add_paragraph('When the script is run directly (__main__), it performs the following startup sequence:')
    for step in [
        'Initialises an OpenTelemetry tracing span (GetConfig) for distributed observability.',
        'Queries POLICY_TABLE DynamoDB for rows where source = aws_config and enabled = True.',
        'Reads the ingest_policy JSON field from the result, which contains: ResourceTypes '
        '(list of resource ID prefixes to query), Annotations (keyword filters for rule '
        'descriptions), Rules (config rule name filters), and ComplianceType (default NON_COMPLIANT).',
        'Initialises an in-memory CSV buffer with the output field headers.',
        'Iterates over each resource type prefix and calls main(item) to fetch results.',
    ]:
        doc.add_paragraph(step, style='List Number')

    # -------------------------------------------------------------------------
    # Section 4: main() — Config Query
    # -------------------------------------------------------------------------
    doc.add_heading('4. The main() Function — Config Aggregator Query', level=1)
    doc.add_paragraph(
        'main(item) builds and executes a SQL query against the AWS Config aggregator '
        'for a given resource ID prefix (e.g. "vol-" for EBS volumes).'
    )

    doc.add_heading('Query Structure', level=2)
    add_code_block(
        doc,
        'SELECT\n'
        '    resourceType, resourceId, resourceName,\n'
        '    configuration.targetResourceType,\n'
        '    configuration.complianceType,\n'
        '    configuration.configRuleList,\n'
        '    configurationItemCaptureTime,\n'
        '    configurationItemStatus,\n'
        '    accountId, awsRegion\n'
        'WHERE configuration.complianceType = \'NON_COMPLIANT\'\n'
        'AND resourceId LIKE \'<item>%\'\n'
        '[AND accountId IN (\'<id1>\', \'<id2>\', ...)]  -- only if ACCOUNT_ID_ALLOWLIST is set\n'
        'ORDER BY accountId DESC'
    )

    doc.add_heading('Production vs. Test Query', level=2)
    doc.add_paragraph(
        'Two versions of the query exist in the code for different use cases:'
    )
    query_headers = ['Version', 'When Active', 'Account Scope']
    query_rows = [
        ['Production query', 'Default — active by default',
         'All accounts, optionally filtered by ACCOUNT_ID_ALLOWLIST env var'],
        ['Test query', 'Commented out — uncomment to activate',
         '12 hardcoded dummy account IDs for isolated testing'],
    ]
    add_styled_table(doc, query_headers, query_rows, col_widths=[1.5, 2.0, 3.0])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('To switch to the test query:').bold = True
    for step in [
        'Comment out the PRODUCTION QUERY block (marked with # TO TEST:).',
        'Uncomment the TEST QUERY block (marked with # TO ACTIVATE:).',
        'Replace the dummy account IDs with real test account IDs.',
        'After testing, reverse the swap to restore production behaviour.',
    ]:
        doc.add_paragraph(step, style='List Number')

    doc.add_heading('Pagination', level=2)
    doc.add_paragraph(
        'The Config API returns results in pages. main() loops using NextToken until all '
        'pages are consumed, accumulating all results into a single list before returning.'
    )

    doc.add_heading('Retry Logic', level=2)
    doc.add_paragraph(
        'On ThrottlingException, the function retries up to 3 times with a linear backoff '
        'of 3 x attempt_number seconds before raising the exception.'
    )

    # -------------------------------------------------------------------------
    # Section 5: Per-Resource Processing
    # -------------------------------------------------------------------------
    doc.add_heading('5. Per-Resource Processing Loop', level=1)
    doc.add_paragraph(
        'After main() returns the full result set, the script iterates over every resource '
        'and applies a multi-step filter chain before writing to the CSV buffer.'
    )

    flow_headers = ['Step', 'Check', 'If True', 'If False']
    flow_rows = [
        ['1', 'Is account in Suspended OU?', 'SKIP resource (log warning)', 'Continue to step 2'],
        ['2', 'Resolve account name (cached)', 'N/A', 'Continue to step 3'],
        ['3', 'Does each config rule match the Rules filter?', 'Fetch annotation', 'Skip that rule'],
        ['4', 'Is account 1.0 legacy? (cached DynamoDB check)', 'SKIP resource (log: 1.0)', 'Continue to step 5'],
        ['5', 'Are any matching annotations found?', 'Write row to CSV buffer', 'SKIP resource (log: no annotations)'],
    ]
    add_styled_table(doc, flow_headers, flow_rows, col_widths=[0.4, 2.2, 1.7, 1.7])

    # -------------------------------------------------------------------------
    # Section 6: Key Functions
    # -------------------------------------------------------------------------
    doc.add_heading('6. Key Functions', level=1)

    func_headers = ['Function', 'Purpose', 'Cache Used']
    func_rows = [
        ['main(item, tries=1)',
         'Executes Config aggregator SQL query for a resource ID prefix. Handles pagination and throttle retries.',
         'None'],
        ['get_rule_description(rule, account, region, ann)',
         'Fetches annotation text for a rule + account + region combination via paginator.',
         'annotation_cache keyed by (rule, account, region)'],
        ['get_account_name(account_id)',
         'Calls organizations:DescribeAccount to resolve account ID to name.',
         'None (raw function)'],
        ['get_account_name_cached(account_id)',
         'Cached wrapper around get_account_name().',
         'account_cache keyed by account_id'],
        ['check_account(account_name)',
         'Queries DynamoDB version table to determine if account is 1.0 legacy.',
         'None (raw function)'],
        ['check_account_cached(account_name)',
         'Cached wrapper around check_account(). Added Feb 17, 2026.',
         'version_cache keyed by account_name'],
        ['get_suspended_account_ids()',
         'Fetches all account IDs under Suspended OU from Organizations. Returns a set.',
         'suspended_account_cache (TTL-based)'],
        ['is_account_in_suspended_ou(account_id)',
         'O(1) membership check against the cached suspended account set.',
         'Uses suspended_account_cache'],
    ]
    add_styled_table(doc, func_headers, func_rows, col_widths=[2.2, 2.8, 1.5])

    # -------------------------------------------------------------------------
    # Section 7: Caching Summary
    # -------------------------------------------------------------------------
    doc.add_heading('7. Caching Summary', level=1)
    doc.add_paragraph(
        'The script maintains four in-process caches to avoid redundant API calls within '
        'a single run. All caches are module-level variables — they live for the lifetime '
        'of the ECS task process and are not shared across runs.'
    )
    cache_headers = ['Cache Variable', 'Function It Backs', 'Keyed By', 'API Call Avoided']
    cache_rows = [
        ['annotation_cache', 'get_rule_description()', '(rule_name, account_id, region)',
         'GetAggregateComplianceDetailsByConfigRule'],
        ['account_cache', 'get_account_name_cached()', 'account_id',
         'organizations:DescribeAccount'],
        ['version_cache', 'check_account_cached()', 'account_name',
         'dynamodb:Query on version table'],
        ['suspended_account_cache', 'get_suspended_account_ids()', 'N/A (full set)',
         'organizations:ListAccountsForParent'],
    ]
    add_styled_table(doc, cache_headers, cache_rows, col_widths=[1.8, 1.8, 1.9, 2.0])

    # -------------------------------------------------------------------------
    # Section 8: Account Exclusion Logic
    # -------------------------------------------------------------------------
    doc.add_heading('8. Account Exclusion Logic', level=1)
    doc.add_paragraph(
        'Two separate exclusion checks prevent certain accounts from appearing in the output. '
        'The Suspended OU check runs first — if an account is suspended, the DynamoDB check '
        'is never reached. This ordering is intentional: the in-memory set lookup (O(1)) is '
        'free, while the DynamoDB call costs money.'
    )
    excl_headers = ['Check', 'When It Runs', 'Source', 'Effect']
    excl_rows = [
        ['Suspended OU check',
         'Per resource, before any other processing',
         'Organizations API (cached set)',
         'Resource skipped, no CSV row written'],
        ['1.0 legacy check',
         'Per resource, after annotation retrieval',
         'DynamoDB version table (cached)',
         'Resource skipped, logged as 1.0'],
    ]
    add_styled_table(doc, excl_headers, excl_rows, col_widths=[1.5, 1.8, 1.8, 1.9])
    doc.add_paragraph(
        'A second Suspended OU check also runs at the S3 upload stage as a '
        'defence-in-depth safety net to prevent any suspended account CSV from reaching S3.'
    )

    # -------------------------------------------------------------------------
    # Section 9: Output
    # -------------------------------------------------------------------------
    doc.add_heading('9. Output — CSV Format and S3 Upload', level=1)
    doc.add_paragraph(
        'After all resources are processed, the in-memory CSV buffer is read into a '
        'pandas DataFrame, grouped by (accountId, accountName), and one CSV file is '
        'uploaded per unique account group.'
    )

    doc.add_heading('CSV Columns', level=2)
    col_headers = ['Column', 'Source']
    col_rows = [
        ['resourceId', 'Config result'],
        ['resourceType', 'Config result'],
        ['resourceName', 'Config result'],
        ['targetResourceType', 'Config result'],
        ['complianceType', 'Config result'],
        ['configRuleName', 'Config result (list of matching rule names)'],
        ['configurationItemCaptureTime', 'Config result'],
        ['configurationItemStatus', 'Config result'],
        ['accountId', 'Config result'],
        ['accountName', 'Organizations API (cached)'],
        ['awsRegion', 'Config result'],
        ['description', 'Config annotation API (cached)'],
    ]
    add_styled_table(doc, col_headers, col_rows, col_widths=[2.5, 4.0])

    doc.add_heading('S3 Key Format', level=2)
    add_code_block(doc, '{BUCKET_PREFIX}/{accountId}-{accountName}_{rule_id}.csv\n\n'
                        'Example: config-reports/daily/111122223333-prod-app-account_1.csv')

    # -------------------------------------------------------------------------
    # Section 10: Error Handling
    # -------------------------------------------------------------------------
    doc.add_heading('10. Error Handling', level=1)
    err_headers = ['Scenario', 'Behaviour']
    err_rows = [
        ['ThrottlingException on Config query',
         'Retry up to 3 times with 3 x attempt second sleep'],
        ['ClientError on DynamoDB query',
         'Log error, return False (account treated as 2.0)'],
        ['Organizations API failure (suspended OU)',
         'Return empty set — no accounts accidentally excluded'],
        ['Empty CSV buffer',
         'pd.errors.EmptyDataError caught and logged'],
        ['S3 upload failure',
         'ClientError caught and logged per account'],
    ]
    add_styled_table(doc, err_headers, err_rows, col_widths=[3.0, 3.5])

    # -------------------------------------------------------------------------
    # Section 11: IAM Permissions
    # -------------------------------------------------------------------------
    doc.add_heading('11. IAM Permissions Required', level=1)
    doc.add_paragraph('The ECS task role must have the following permissions:')
    iam_headers = ['Permission', 'Used By']
    iam_rows = [
        ['config:SelectAggregateResourceConfig', 'main() — Config aggregator query'],
        ['config:GetAggregateComplianceDetailsByConfigRule', 'get_rule_description()'],
        ['dynamodb:Query', 'check_account() — version table lookup'],
        ['dynamodb:Query', '__main__ — policy table lookup'],
        ['organizations:DescribeAccount', 'get_account_name()'],
        ['organizations:ListAccountsForParent', 'get_suspended_account_ids()'],
        ['s3:PutObject', 'S3 CSV upload'],
    ]
    add_styled_table(doc, iam_headers, iam_rows, col_widths=[3.5, 3.0])

    # -------------------------------------------------------------------------
    # Save
    # -------------------------------------------------------------------------
    output_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        'config_aggregator_how_it_works.docx'
    )
    doc.save(output_path)
    print(f'Word document saved to: {output_path}')


if __name__ == '__main__':
    main()
